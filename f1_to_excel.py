#!/usr/bin/env python3
"""
f1_to_excel.py - Pull F1 race data from FastF1 and write one Excel workbook per season.

Scope: race sessions only (no sprints, no practice/quali), DRY races only by default.
A race counts as wet/mixed if any lap was run on INTERMEDIATE or WET tyres, or if
rainfall was recorded while the race was running. Excluded races still appear in the
"Races" sheet with the reason, so nothing disappears silently.

Usage:
    pip install fastf1 openpyxl pandas
    python f1_to_excel.py --years 2024                    # one season
    python f1_to_excel.py --years 2022 2023 2024 2025     # several seasons
    python f1_to_excel.py --years 2024 --rounds 1 2 3     # quick test on 3 races
    python f1_to_excel.py --years 2019 --include-wet      # keep wet races too

Seasons outside 2021-2022 need an F1TV token (paid tier): python -m fastf1 auth f1tv --authenticate

Output: <out>/F1_<year>_dry_races.xlsx
FastF1 caches everything it downloads in <cache>, so re-running is fast and a run
that stops halfway (rate limit, network) resumes where it left off.
"""
import argparse
import logging
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import numpy as np
import pandas as pd

WET_COMPOUNDS = {"INTERMEDIATE", "WET"}
WEATHER_COLS = ["AirTemp", "TrackTemp", "Humidity", "Pressure", "Rainfall", "WindSpeed", "WindDirection"]

log = logging.getLogger("f1_to_excel")


# --------------------------------------------------------------------------- helpers
def to_seconds(s: pd.Series) -> pd.Series:
    """Timedelta column -> float seconds (Excel can't store timedeltas cleanly)."""
    return pd.to_timedelta(s).dt.total_seconds().round(3)


def status_has(ts: pd.Series, code: str) -> pd.Series:
    """FastF1 lap TrackStatus is a string of every status code seen during the lap, e.g. '1', '24', '4'."""
    return ts.str.contains(code, regex=False)


# --------------------------------------------------------------------------- per-race processing
def build_laps(laps: pd.DataFrame, meta: dict) -> pd.DataFrame:
    df = pd.DataFrame({
        "Year": meta["Year"],
        "Round": meta["Round"],
        "Event": meta["Event"],
        "Driver": laps["Driver"].values,
        "DriverNumber": laps["DriverNumber"].values,
        "Team": laps["Team"].values,
        "LapNumber": laps["LapNumber"].astype("Int64").values,
        "Position": laps["Position"].astype("Int64").values,
        "LapTime_s": to_seconds(laps["LapTime"]).values,
        "Sector1_s": to_seconds(laps["Sector1Time"]).values,
        "Sector2_s": to_seconds(laps["Sector2Time"]).values,
        "Sector3_s": to_seconds(laps["Sector3Time"]).values,
        "Compound": laps["Compound"].values,
        "TyreLife": laps["TyreLife"].astype("Int64").values,
        "FreshTyre": laps["FreshTyre"].values,
        "Stint": laps["Stint"].astype("Int64").values,
        "PitInTime_s": to_seconds(laps["PitInTime"]).values,
        "PitOutTime_s": to_seconds(laps["PitOutTime"]).values,
        "LapStartTime_s": to_seconds(laps["LapStartTime"]).values,
        "LapEndTime_s": to_seconds(laps["Time"]).values,
        "SpeedST_kph": pd.to_numeric(laps["SpeedST"], errors="coerce").values,
        "TrackStatus": laps["TrackStatus"].fillna("").astype(str).values,
        "Deleted": laps["Deleted"].fillna(False).astype(bool).values,
        "IsAccurate": laps["IsAccurate"].fillna(False).astype(bool).values,
    })
    ts = df["TrackStatus"]
    df["IsInLap"] = df["PitInTime_s"].notna()
    df["IsOutLap"] = df["PitOutTime_s"].notna() & (df["LapNumber"] > 1)
    df["GreenFlag"] = ts == "1"
    df["Yellow"] = status_has(ts, "2")
    df["SC"] = status_has(ts, "4")
    df["VSC"] = status_has(ts, "6") | status_has(ts, "7")
    df["RedFlag"] = status_has(ts, "5")
    # Clean lap = usable for degradation / pace fitting.
    df["CleanLap"] = (
        df["LapTime_s"].notna()
        & df["GreenFlag"]
        & ~df["IsInLap"]
        & ~df["IsOutLap"]
        & df["IsAccurate"]
        & ~df["Deleted"]
        & (df["LapNumber"] > 1)
        & ~df["Compound"].isin(WET_COMPOUNDS)
    )
    return df


def build_stints(lapdf: pd.DataFrame) -> pd.DataFrame:
    d = lapdf.dropna(subset=["Stint"])
    if d.empty:
        return pd.DataFrame()
    g = d.groupby(["Year", "Round", "Event", "Driver", "Team", "Stint"], sort=False)
    out = g.agg(
        Compound=("Compound", "first"),
        StartLap=("LapNumber", "min"),
        EndLap=("LapNumber", "max"),
        StintLaps=("LapNumber", "count"),
        TyreLifeAtStart=("TyreLife", "first"),
        FreshTyre=("FreshTyre", "first"),
        CleanLaps=("CleanLap", "sum"),
    ).reset_index()
    clean = d[d["CleanLap"]].groupby(["Driver", "Stint"])["LapTime_s"]
    med = clean.median().rename("MedianCleanLapTime_s")
    best = clean.min().rename("BestCleanLapTime_s")
    out = out.join(med, on=["Driver", "Stint"]).join(best, on=["Driver", "Stint"])
    return out.sort_values(["Driver", "Stint"]).reset_index(drop=True)


def build_pitstops(lapdf: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for drv, dl in lapdf.groupby("Driver", sort=False):
        dl = dl.set_index("LapNumber")
        stop_no = 0
        for lap_no, r in dl[dl["IsInLap"]].iterrows():
            stop_no += 1
            nxt = dl.loc[lap_no + 1] if (lap_no + 1) in dl.index else None
            has_out = nxt is not None and pd.notna(nxt["PitOutTime_s"])
            rows.append({
                "Year": r["Year"], "Round": r["Round"], "Event": r["Event"],
                "Driver": drv, "Team": r["Team"], "StopNumber": stop_no, "InLap": int(lap_no),
                "CompoundBefore": r["Compound"],
                "CompoundAfter": nxt["Compound"] if has_out else None,
                "PitLaneTime_s": round(nxt["PitOutTime_s"] - r["PitInTime_s"], 3) if has_out else np.nan,
                "InLapTime_s": r["LapTime_s"],
                "OutLapTime_s": nxt["LapTime_s"] if has_out else np.nan,
                "TrackStatusInLap": r["TrackStatus"],
                "UnderSCorVSC": bool(r["SC"] or r["VSC"]),
                "UnderRedFlag": bool(r["RedFlag"] or (has_out and nxt["RedFlag"])),
                "Note": "" if has_out else "No out-lap (retired in pits or race ended)",
            })
    return pd.DataFrame(rows)


def race_window(laps: pd.DataFrame):
    return laps["LapStartTime"].min(), laps["Time"].max()


def classify_conditions(laps: pd.DataFrame, weather: pd.DataFrame):
    start, end = race_window(laps)
    wet_laps = int(laps["Compound"].isin(WET_COMPOUNDS).sum())
    rain = False
    if weather is not None and not weather.empty and "Rainfall" in weather:
        w = weather[(weather["Time"] >= start) & (weather["Time"] <= end)]
        rain = bool(w["Rainfall"].fillna(False).astype(bool).any())
    reasons = []
    if wet_laps:
        reasons.append(f"{wet_laps} laps on INTERMEDIATE/WET")
    if rain:
        reasons.append("rainfall recorded during race")
    return ("WET/MIXED" if reasons else "DRY"), wet_laps, rain, "; ".join(reasons)


def build_weather(weather: pd.DataFrame, laps: pd.DataFrame, meta: dict) -> pd.DataFrame:
    if weather is None or weather.empty:
        return pd.DataFrame()
    start, end = race_window(laps)
    w = weather[(weather["Time"] >= start) & (weather["Time"] <= end)].copy()
    out = pd.DataFrame({"Year": meta["Year"], "Round": meta["Round"], "Event": meta["Event"],
                        "SessionTime_s": to_seconds(w["Time"]).values})
    for c in WEATHER_COLS:
        if c in w:
            out[c] = w[c].values
    return out


def build_track_status(track_status: pd.DataFrame, meta: dict) -> pd.DataFrame:
    if track_status is None or track_status.empty:
        return pd.DataFrame()
    return pd.DataFrame({
        "Year": meta["Year"], "Round": meta["Round"], "Event": meta["Event"],
        "SessionTime_s": to_seconds(track_status["Time"]).values,
        "Status": track_status["Status"].astype(str).values,
        "Message": track_status["Message"].values,
    })


def build_results(results: pd.DataFrame, meta: dict) -> pd.DataFrame:
    if results is None or results.empty:
        return pd.DataFrame()
    cols = ["Position", "ClassifiedPosition", "GridPosition", "Abbreviation", "FullName",
            "TeamName", "Status", "Points", "Laps"]
    out = results[[c for c in cols if c in results]].copy()
    if "Time" in results:
        out["ResultTime_s"] = to_seconds(results["Time"])
    out.insert(0, "Event", meta["Event"])
    out.insert(0, "Round", meta["Round"])
    out.insert(0, "Year", meta["Year"])
    return out.reset_index(drop=True)


def summarize_race(lapdf, pits, tstat, wdf, results, meta) -> dict:
    ts = tstat["Status"] if not tstat.empty else pd.Series(dtype=str)
    green_pits = pits[~pits["UnderRedFlag"] & pits["PitLaneTime_s"].notna()] if not pits.empty else pits
    classified = results[pd.to_numeric(results.get("ClassifiedPosition"), errors="coerce").notna()] \
        if not results.empty and "ClassifiedPosition" in results else pd.DataFrame()
    if not classified.empty and not pits.empty:
        stops = pits.groupby("Driver").size().reindex(classified["Abbreviation"]).fillna(0)
        mean_stops = round(float(stops.mean()), 2)
    else:
        mean_stops = np.nan
    return {
        "Year": meta["Year"], "Round": meta["Round"], "Event": meta["Event"],
        "RaceLaps": int(lapdf["LapNumber"].max()),
        "Drivers": int(lapdf["Driver"].nunique()),
        "Classified": len(classified) if not classified.empty else np.nan,
        "MeanStopsPerClassifiedDriver": mean_stops,
        "MedianPitLaneTime_s": round(float(green_pits["PitLaneTime_s"].median()), 3) if not green_pits.empty else np.nan,
        "StopsUnderSCorVSC": int(pits["UnderSCorVSC"].sum()) if not pits.empty else 0,
        "SCDeployments": int((ts == "4").sum()),
        "VSCDeployments": int((ts == "6").sum()),
        "RedFlags": int((ts == "5").sum()),
        "CleanLapShare": round(float(lapdf["CleanLap"].mean()), 3),
        "CompoundsUsed": ", ".join(sorted(lapdf["Compound"].dropna().astype(str).unique())),
        "MeanAirTemp_C": round(float(wdf["AirTemp"].mean()), 1) if not wdf.empty else np.nan,
        "MeanTrackTemp_C": round(float(wdf["TrackTemp"].mean()), 1) if not wdf.empty else np.nan,
    }


def process_race(laps, results, weather, track_status, meta, include_wet):
    """Pure-DataFrame core: everything except the network call. Returns (race_row, tables|None)."""
    condition, wet_laps, rain, reason = classify_conditions(laps, weather)
    race_row = {**meta, "Condition": condition, "WetLapCount": wet_laps,
                "RainfallDuringRace": rain, "Included": True, "ExcludeReason": "", "LoadStatus": "OK"}
    if condition != "DRY" and not include_wet:
        race_row.update(Included=False, ExcludeReason=reason)
        return race_row, None
    lapdf = build_laps(laps, meta)
    pits = build_pitstops(lapdf)
    wdf = build_weather(weather, laps, meta)
    tstat = build_track_status(track_status, meta)
    res = build_results(results, meta)
    tables = {
        "Laps": lapdf,
        "Stints": build_stints(lapdf),
        "PitStops": pits,
        "TrackStatus": tstat,
        "Weather": wdf,
        "Results": res,
        "TrackSummary": pd.DataFrame([summarize_race(lapdf, pits, tstat, wdf, res, meta)]),
    }
    return race_row, tables


# --------------------------------------------------------------------------- F1TV auth
F1_ARCHIVE_HOST = "livetiming.formula1.com"


class _F1TVAuth:
    """Adds the F1TV token to requests for F1's own timing archive only (never the mirror or Ergast)."""
    def __init__(self, token):
        self.token = token

    def __call__(self, r):
        if urlparse(r.url).hostname == F1_ARCHIVE_HOST:
            r.headers["Authorization"] = f"Bearer {self.token}"
        return r


def attach_f1tv_token():
    """The timing archive now returns 401 without an F1TV login. FastF1 3.8 only sends the
    token for live timing, so attach it to FastF1's archive requests ourselves."""
    import fastf1
    import jwt
    from fastf1.internals.f1auth import AUTH_DATA_FILE
    token = AUTH_DATA_FILE.read_text().strip() if AUTH_DATA_FILE.exists() else ""
    if not token:
        log.warning("No F1TV token found; only seasons on the free mirror (2021-2022) will load. "
                    "Run: python -m fastf1 auth f1tv --authenticate")
        return
    exp = jwt.decode(token, options={"verify_signature": False}).get("exp")
    if exp and exp < datetime.now(timezone.utc).timestamp():
        log.warning("F1TV token expired %s UTC; re-authenticate or archive downloads will fail.",
                    datetime.fromtimestamp(exp, timezone.utc).strftime("%Y-%m-%d %H:%M"))
    for sess in (fastf1.Cache._requests_session, fastf1.Cache._requests_session_cached):
        if sess is not None:
            sess.auth = _F1TVAuth(token)


# --------------------------------------------------------------------------- FastF1 I/O
def load_season(year, rounds, include_wet, cache_dir, offline=False):
    import fastf1
    fastf1.Cache.enable_cache(str(cache_dir))
    if offline:
        # Cache only: no network requests, so no F1TV token is needed. Uncached races show FAILED.
        fastf1.Cache.offline_mode(True)
    else:
        attach_f1tv_token()
    schedule = fastf1.get_event_schedule(year, include_testing=False)
    if rounds:
        schedule = schedule[schedule["RoundNumber"].isin(rounds)]
    today = pd.Timestamp(datetime.now(timezone.utc).date())

    race_rows, collected = [], {}
    for _, ev in schedule.iterrows():
        meta = {"Year": year, "Round": int(ev["RoundNumber"]), "Event": ev["EventName"],
                "Country": ev["Country"], "Location": ev["Location"],
                "Date": pd.Timestamp(ev["EventDate"]).date()}
        if pd.Timestamp(ev["EventDate"]).normalize() >= today:
            race_rows.append({**meta, "Included": False, "LoadStatus": "NOT RUN YET"})
            continue
        log.info("%d R%02d %s", year, meta["Round"], meta["Event"])
        try:
            s = fastf1.get_session(year, meta["Round"], "R")
            s.load(laps=True, telemetry=False, weather=True, messages=True)
            if s.laps is None or s.laps.empty:
                raise RuntimeError("no lap data returned")
            row, tables = process_race(s.laps, s.results, s.weather_data, s.track_status, meta, include_wet)
        except Exception as e:  # keep going; record the failure in the Races sheet
            log.warning("  failed: %s", e)
            race_rows.append({**meta, "Included": False, "LoadStatus": f"FAILED: {type(e).__name__}: {e}"[:250]})
            continue
        race_rows.append(row)
        if tables:
            for k, v in tables.items():
                if not v.empty:
                    collected.setdefault(k, []).append(v)
    sheets = {k: pd.concat(v, ignore_index=True) for k, v in collected.items()}
    return pd.DataFrame(race_rows), sheets


# --------------------------------------------------------------------------- Excel output
# Data dictionary: one entry per column. Fields: (sheet, column, unit, meaning, how to use it).
# Sheet "All" = the column means the same thing in every sheet it appears in.
# When you add or rename a column, add it here in the same change; check_dictionary() warns on gaps.
DICTIONARY = [
    # ---- columns shared by several sheets
    ("All", "Year", "-", "Season.", "Join key. Year + Round = one race."),
    ("All", "Round", "-", "Round number in that season's calendar (1 = first race of the year).", "Join key. Use Year + Round, not Event, to join sheets."),
    ("All", "Event", "-", "Official Grand Prix name, e.g. 'Italian Grand Prix'.", "Readable label only."),
    ("All", "Driver", "-", "Three-letter driver code, e.g. VER = Verstappen, HAM = Hamilton (full names in Results.FullName).", "Join key. Year + Round + Driver = one driver's race."),
    ("All", "Team", "-", "Constructor (team) name.", "Car pace differs by team; pace and degradation are fitted per team or car."),
    ("All", "SessionTime_s", "s", "Seconds since the session clock started. The clock starts before lights out, so this is NOT race time.", "Lines up events across sheets: track status changes, weather samples and lap end times share this clock."),
    ("All", "Compound", "-", "Tyre compound: SOFT, MEDIUM or HARD (dry), INTERMEDIATE or WET. See the Legend sheet: the names are relative to each race weekend.", "Pace and wear are fitted per compound."),
    ("All", "FreshTyre", "True/False", "True if the tyre set was new (never used, not even in qualifying) when fitted.", "Used sets start the stint with some wear already."),
    ("All", "Stint", "-", "Stint number: 1 = tyres the driver started on, 2 = after the first tyre change, and so on.", "Number of stints - 1 = number of pit stops."),

    # ---- Races: one row per event in the season's calendar
    ("Races", "Country", "-", "Country of the event.", ""),
    ("Races", "Location", "-", "City or circuit name.", ""),
    ("Races", "Date", "date", "Race day.", ""),
    ("Races", "Condition", "-", "DRY, or WET/MIXED if any lap used INTERMEDIATE/WET tyres or rain was recorded during the race.", "Project scope is DRY races only. WET/MIXED races are excluded from the other sheets."),
    ("Races", "WetLapCount", "laps", "Total laps (all drivers together) driven on INTERMEDIATE or WET tyres.", "Reason for exclusion."),
    ("Races", "RainfallDuringRace", "True/False", "True if the official weather feed reported rain at any point between race start and finish.", "Reason for exclusion."),
    ("Races", "Included", "True/False", "True if this race's data is in the other sheets.", "False = excluded (see ExcludeReason) or not loaded (see LoadStatus)."),
    ("Races", "ExcludeReason", "-", "Why a loaded race was left out, e.g. '12 laps on INTERMEDIATE/WET'.", ""),
    ("Races", "LoadStatus", "-", "OK; NOT RUN YET (race is in the future); or FAILED with the download error.", "FAILED races can be retried by re-running the script."),

    # ---- TrackSummary: one row per included race
    ("TrackSummary", "RaceLaps", "laps", "Laps completed by the winner (the race distance, unless the race was shortened).", ""),
    ("TrackSummary", "Drivers", "drivers", "Drivers who completed at least one lap.", ""),
    ("TrackSummary", "Classified", "drivers", "Drivers officially classified at the finish (blank if results data is missing).", ""),
    ("TrackSummary", "MeanStopsPerClassifiedDriver", "stops", "Average number of pit stops made by classified drivers.", "About 1 = mostly one-stop race, about 2 = two-stop. Shows how strategy-heavy the track is."),
    ("TrackSummary", "MedianPitLaneTime_s", "s", "Median time from pit entry to pit exit, excluding stops under a red flag.", "Part of the pit loss. Full pit loss also includes slowing for the in-lap and accelerating on the out-lap (see PitStops)."),
    ("TrackSummary", "StopsUnderSCorVSC", "stops", "Pit stops made on a lap with a safety car or virtual safety car.", "These stops cost less time, so they distort average pit loss."),
    ("TrackSummary", "SCDeployments", "count", "Times the safety car was deployed.", "Safety-car frequency per track feeds the simulator's random-event model."),
    ("TrackSummary", "VSCDeployments", "count", "Times the virtual safety car was deployed.", "As above, for the VSC."),
    ("TrackSummary", "RedFlags", "count", "Times the race was stopped with a red flag.", "Drivers can change tyres for free under a red flag."),
    ("TrackSummary", "CleanLapShare", "0-1", "Share of all laps that are clean laps (0.85 = 85%).", "How much usable pace data the race gives."),
    ("TrackSummary", "CompoundsUsed", "-", "Compounds that appear in the race.", ""),
    ("TrackSummary", "MeanAirTemp_C", "degC", "Average air temperature during the race.", "Hot races usually mean more tyre wear."),
    ("TrackSummary", "MeanTrackTemp_C", "degC", "Average track surface temperature during the race.", "As above; track temperature matters more than air temperature."),

    # ---- Laps: one row per lap per driver
    ("Laps", "DriverNumber", "-", "Car number.", ""),
    ("Laps", "LapNumber", "lap", "Lap of the race, starting at 1.", ""),
    ("Laps", "Position", "1-20", "Race position at the end of this lap (1 = leading).", "Blank in rare cases where timing has no position (e.g. the lap a car retires)."),
    ("Laps", "LapTime_s", "s", "Time to complete the lap. Blank if timing did not record it.", "Raw lap time. Use only CleanLap = True laps to fit pace and wear."),
    ("Laps", "Sector1_s", "s", "Time for the first of the three timing sectors of the lap.", ""),
    ("Laps", "Sector2_s", "s", "Time for the second sector.", ""),
    ("Laps", "Sector3_s", "s", "Time for the third sector.", ""),
    ("Laps", "TyreLife", "laps", "Laps driven on this tyre set, counting this lap, plus any laps the set did earlier in the weekend.", "Tyre age: the main input for tyre wear."),
    ("Laps", "PitInTime_s", "s (session clock)", "When the car entered the pit lane. Only filled on in-laps.", ""),
    ("Laps", "PitOutTime_s", "s (session clock)", "When the car left the pit lane. Only filled on out-laps (and lap 1 for pit-lane starters).", ""),
    ("Laps", "LapStartTime_s", "s (session clock)", "When the lap started.", ""),
    ("Laps", "LapEndTime_s", "s (session clock)", "When the lap ended.", "Gap to the car ahead after lap N = difference in LapEndTime_s between the two drivers on lap N."),
    ("Laps", "SpeedST_kph", "km/h", "Speed at the speed trap on this lap.", "A driver's maximum approximates top speed (TUM's vel_max)."),
    ("Laps", "TrackStatus", "codes", "Every track status code seen during the lap, joined together. '1' = all green; '124' = green, then yellow, then safety car. Excel may show it as a number: read it digit by digit. Codes in the Legend sheet.", "Used to flag safety-car and yellow-flag laps."),
    ("Laps", "Deleted", "True/False", "True if the stewards deleted this lap time (usually track limits).", ""),
    ("Laps", "IsAccurate", "True/False", "FastF1's check that the lap's timing data is consistent and trustworthy.", "False laps are left out of clean laps."),
    ("Laps", "IsInLap", "True/False", "True if the driver entered the pits at the end of this lap.", "In-laps are slower and are not clean laps."),
    ("Laps", "IsOutLap", "True/False", "True if this lap started from the pit exit (lap 1 is never counted).", "Out-laps are slower (cold tyres, pit lane) and are not clean laps."),
    ("Laps", "GreenFlag", "True/False", "True if the whole lap was under green flag (TrackStatus is exactly '1').", ""),
    ("Laps", "Yellow", "True/False", "True if there was a yellow flag somewhere on track during the lap.", ""),
    ("Laps", "SC", "True/False", "True if the safety car was out at any point during the lap.", "Pitting under the safety car is cheap: the field is slow, so less time is lost."),
    ("Laps", "VSC", "True/False", "True if a virtual safety car was out (or ending) during the lap.", "Same idea as SC, smaller effect."),
    ("Laps", "RedFlag", "True/False", "True if the race was red-flagged during the lap.", ""),
    ("Laps", "CleanLap", "True/False", "True if the lap is representative racing pace: green flag, not lap 1, not an in- or out-lap, accurate, not deleted, dry tyres.", "Only these laps are used to fit pace and tyre wear."),

    # ---- Stints: one row per stint (one tyre set) per driver
    ("Stints", "StartLap", "lap", "First lap of the stint.", ""),
    ("Stints", "EndLap", "lap", "Last lap of the stint (the in-lap, or the final lap).", ""),
    ("Stints", "StintLaps", "laps", "Laps recorded in the stint.", "Stint length: how long a set lasted."),
    ("Stints", "TyreLifeAtStart", "laps", "Tyre age on the first lap of the stint. 1 = new set; more than 1 = used set.", ""),
    ("Stints", "CleanLaps", "laps", "Clean laps in the stint.", "Few clean laps = an unreliable wear estimate for this stint."),
    ("Stints", "MedianCleanLapTime_s", "s", "Median lap time of the stint's clean laps. NOT fuel-corrected.", "Cars get lighter as fuel burns, which hides some tyre wear. Correct for fuel before fitting wear."),
    ("Stints", "BestCleanLapTime_s", "s", "Fastest clean lap of the stint.", ""),

    # ---- PitStops: one row per pit stop
    ("PitStops", "StopNumber", "-", "1 = the driver's first stop of the race, 2 = second, and so on.", ""),
    ("PitStops", "InLap", "lap", "Lap at the end of which the driver pitted.", "This is the decision the project models: which lap to pit on."),
    ("PitStops", "CompoundBefore", "-", "Compound on the car before the stop.", ""),
    ("PitStops", "CompoundAfter", "-", "Compound fitted at the stop. Blank if there was no out-lap.", ""),
    ("PitStops", "PitLaneTime_s", "s", "Pit entry to pit exit: driving the pit lane at the speed limit plus standing still.", "NOT the full pit loss. Pit loss = (InLapTime_s + OutLapTime_s) - 2 x a normal clean lap."),
    ("PitStops", "InLapTime_s", "s", "Lap time of the in-lap.", "Used to compute pit loss."),
    ("PitStops", "OutLapTime_s", "s", "Lap time of the out-lap.", "Used to compute pit loss."),
    ("PitStops", "TrackStatusInLap", "codes", "Track status codes during the in-lap (same format as Laps.TrackStatus).", ""),
    ("PitStops", "UnderSCorVSC", "True/False", "True if the stop was made on a safety-car or VSC lap.", "Cheap stops. Kept, but modelled separately from green-flag stops."),
    ("PitStops", "UnderRedFlag", "True/False", "True if the tyre change happened during a red flag.", "Free tyre change; left out of pit-loss averages."),
    ("PitStops", "Note", "-", "Data caveat, e.g. no out-lap because the car retired in the pits.", ""),

    # ---- TrackStatus: one row per change of track status
    ("TrackStatus", "Status", "code", "New track status code from this moment on. Codes in the Legend sheet.", ""),
    ("TrackStatus", "Message", "-", "Text for the code: AllClear, Yellow, SCDeployed, Red, VSCDeployed, VSCEnding.", ""),

    # ---- Weather: about one sample per minute during the race
    ("Weather", "AirTemp", "degC", "Air temperature.", ""),
    ("Weather", "TrackTemp", "degC", "Track surface temperature.", "Higher track temperature usually means more tyre wear."),
    ("Weather", "Humidity", "%", "Relative humidity.", ""),
    ("Weather", "Pressure", "mbar", "Air pressure.", ""),
    ("Weather", "Rainfall", "True/False", "True if rain was falling at this sample.", "Used to classify races as DRY or WET/MIXED."),
    ("Weather", "WindSpeed", "m/s", "Wind speed.", ""),
    ("Weather", "WindDirection", "degrees", "Wind direction, 0-359 (0 = north, 90 = east).", ""),

    # ---- Results: one row per driver per race (official classification)
    ("Results", "Position", "1-20", "Finishing position. Retired drivers are also given a number here.", "Use ClassifiedPosition to tell finishers from retirements."),
    ("Results", "ClassifiedPosition", "-", "Official result: a number, or a letter for drivers not classified (R retired, D disqualified, etc.; see Legend).", "The simulator's finishing order is compared against this."),
    ("Results", "GridPosition", "1-20", "Starting position. 0 = started from the pit lane.", ""),
    ("Results", "Abbreviation", "-", "Three-letter driver code (same as Driver in other sheets).", "Join key with Driver."),
    ("Results", "FullName", "-", "Driver's full name.", ""),
    ("Results", "TeamName", "-", "Team name.", ""),
    ("Results", "Status", "-", "Finished; '+1 Lap' or 'Lapped' = finished one or more laps behind the winner; anything else = retirement reason (e.g. Engine, Collision).", ""),
    ("Results", "Points", "points", "Championship points scored.", ""),
    ("Results", "Laps", "laps", "Laps completed.", ""),
    ("Results", "ResultTime_s", "s", "Winner: total race time. Other drivers on the winner's lap: gap behind the winner. Lapped drivers: not a usable gap (ignore it).", "Compare race times only for drivers on the lead lap."),

    # ---- this file's own sheets
    ("DataDictionary", "Unit", "-", "Unit of the column. s = seconds; s (session clock) = seconds on the session clock; True/False = yes/no flag; '-' = text or ID.", ""),
    ("DataDictionary", "How it is used", "-", "Why the column matters for the pit-stop model, or a caution when reading it.", ""),
]

# Code lists for the Legend sheet: (group, code, meaning).
LEGEND = [
    ("How to read codes", "TrackStatus", "Laps.TrackStatus lists every code seen during the lap, joined: '1' = all green; '12' = green then yellow; '124' = green, yellow, then safety car. Read it digit by digit."),
    ("TrackStatus code", "1", "Track clear (green)."),
    ("TrackStatus code", "2", "Yellow flag somewhere on track: no overtaking in that zone."),
    ("TrackStatus code", "4", "Safety car deployed: field bunches up behind it at reduced speed."),
    ("TrackStatus code", "5", "Red flag: race stopped, cars go to the pit lane and may change tyres for free."),
    ("TrackStatus code", "6", "Virtual safety car (VSC) deployed: every driver must slow to a set lap delta; gaps stay roughly the same."),
    ("TrackStatus code", "7", "VSC ending: racing resumes shortly."),
    ("Compound", "SOFT / MEDIUM / HARD", "Dry tyres. For each race Pirelli picks three of its compounds (C1 = hardest ... C5/C6 = softest) and names them HARD, MEDIUM, SOFT. So SOFT at one track can be a different rubber from SOFT at another: compare compounds within a race, not across tracks."),
    ("Compound", "Faster vs longer", "SOFT is fastest but wears quickest; HARD is slowest but lasts longest."),
    ("Compound", "Two-compound rule", "In a dry race each driver must use at least two different dry compounds, so at least one pit stop is required."),
    ("Compound", "INTERMEDIATE / WET", "Rain tyres. Any lap on these makes the race WET/MIXED and it is excluded."),
    ("Compound", "UNKNOWN / blank", "Timing did not record the compound."),
    ("Races.Condition", "DRY", "No laps on rain tyres and no rainfall recorded during the race. In scope."),
    ("Races.Condition", "WET/MIXED", "Rain tyres used or rainfall recorded. Out of scope; excluded."),
    ("Races.LoadStatus", "OK", "Race downloaded and processed."),
    ("Races.LoadStatus", "NOT RUN YET", "Race is later in the season."),
    ("Races.LoadStatus", "FAILED: ...", "Download failed (usually a rate limit or missing login). Re-run the script."),
    ("Results.ClassifiedPosition", "number", "Official finishing position."),
    ("Results.ClassifiedPosition", "R", "Retired."),
    ("Results.ClassifiedPosition", "D", "Disqualified."),
    ("Results.ClassifiedPosition", "E", "Excluded."),
    ("Results.ClassifiedPosition", "W", "Withdrawn (did not start)."),
    ("Results.ClassifiedPosition", "F", "Failed to qualify."),
    ("Results.ClassifiedPosition", "N", "Not classified (finished but completed too little of the race distance)."),
    ("Results.GridPosition", "0", "Started from the pit lane."),
    ("Terms", "Stint", "The laps driven on one set of tyres, between two pit stops."),
    ("Terms", "In-lap / out-lap", "The lap into the pits and the lap back out of them. Both are slower than normal laps."),
    ("Terms", "Pit loss", "Total time a stop costs compared with staying out: about (in-lap + out-lap) - 2 x a normal lap. Larger than PitLaneTime_s."),
    ("Terms", "Clean lap", "A lap that shows true racing pace: green flag, no pit stop, not lap 1, accurate timing, dry tyres."),
    ("Terms", "Degradation (wear)", "How much slower a tyre gets each lap as it wears, in seconds per lap."),
    ("Terms", "Session clock", "Columns ending in _s marked 'session clock' count from when timing started, not from lights out."),
    ("Cells", "blank", "Not recorded, or does not apply (e.g. PitInTime_s on a lap with no stop)."),
    ("Cells", "True / False", "Yes / no flag."),
]

# What one row means in each sheet (shown in README and as the first lines of each sheet's guide).
SHEET_GUIDE = [
    ("DataDictionary", "Every column in this file: unit, meaning and how it is used. Also shown when you hover over a column header."),
    ("Legend", "Codes and terms: track status codes, tyre compounds, result letters."),
    ("Races", "One row per race in the season's calendar, including excluded and future races, with the reason."),
    ("TrackSummary", "One row per included race: laps, stops, safety cars, pit-lane time, temperatures. Start here to compare tracks."),
    ("Laps", "One row per lap per driver. The raw data everything else is built from."),
    ("Stints", "One row per stint (one set of tyres) per driver."),
    ("PitStops", "One row per pit stop."),
    ("TrackStatus", "One row each time the track status changed (green, yellow, safety car, VSC, red flag)."),
    ("Weather", "About one row per minute of the race."),
    ("Results", "One row per driver per race: the official result."),
]
SHEET_ORDER = ["README", "DataDictionary", "Legend", "Races", "TrackSummary", "Laps", "Stints",
               "PitStops", "TrackStatus", "Weather", "Results"]


def dictionary_lookup():
    """(sheet, column) -> (unit, meaning, use). Sheet-specific entries win over 'All'."""
    shared = {col: (u, m, h) for s, col, u, m, h in DICTIONARY if s == "All"}
    specific = {(s, col): (u, m, h) for s, col, u, m, h in DICTIONARY if s != "All"}
    return lambda sheet, col: specific.get((sheet, col)) or shared.get(col)


def check_dictionary(sheets):
    """Warn about any column that has no dictionary entry (keeps the legend complete)."""
    lookup = dictionary_lookup()
    missing = [f"{name}.{col}" for name, df in sheets.items()
               if name not in ("README", "DataDictionary", "Legend")
               for col in df.columns if lookup(name, col) is None]
    if missing:
        log.warning("Columns missing from DICTIONARY (add them): %s", ", ".join(missing))
    return missing


def dictionary_frame(sheets):
    """Dictionary rows for the columns actually in this workbook, in sheet order."""
    lookup = dictionary_lookup()
    rows = []
    for name in SHEET_ORDER:
        df = sheets.get(name)
        if df is None or not len(df.columns) or name in ("README", "DataDictionary", "Legend"):
            continue
        for col in df.columns:
            unit, meaning, use = lookup(name, col) or ("", "(not documented)", "")
            rows.append((name, col, unit, meaning, use))
    return pd.DataFrame(rows, columns=["Sheet", "Column", "Unit", "Meaning", "How it is used"])


def readme_frame(year, include_wet, races, generated=None):
    import importlib.metadata as md
    try:
        ff1v = md.version("fastf1")
    except Exception:
        ff1v = "unknown"
    n_inc = int(races["Included"].fillna(False).sum()) if not races.empty else 0
    items = [
        ("Season", year),
        ("Generated (UTC)", generated or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Source", f"Official F1 timing data via the FastF1 Python library ({ff1v})"),
        ("Session type", "Grand Prix race only (sprints, practice and qualifying excluded)"),
        ("Filter", "All races" if include_wet else "Dry races only (see Races sheet for exclusions and reasons)"),
        ("Races included", n_inc),
        ("Races in schedule", len(races)),
        ("Units", "Times in seconds (s), speeds in km/h, temperatures in degC. Every column's unit is in DataDictionary."),
        ("Column meanings", "Hover over any column header to see what it means, or open the DataDictionary sheet."),
        ("How sheets connect", "Year + Round identify a race; Year + Round + Driver identify one driver's race. Use these to join sheets."),
    ]
    items += [(f"Sheet: {name}", text) for name, text in SHEET_GUIDE]
    items.append(("Notes", "Generated by f1_to_excel.py. Do not edit by hand; re-run the script instead."))
    return pd.DataFrame(items, columns=["Item", "Value"])


def write_workbook(path, sheets):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name in SHEET_ORDER:
            if name in sheets and not sheets[name].empty:
                sheets[name].to_excel(xw, sheet_name=name, index=False)
        style_workbook(xw.book)


def style_workbook(wb):
    """Header colours, hover notes from DICTIONARY, filters, column widths."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    lookup = dictionary_lookup()
    docs = ("README", "DataDictionary", "Legend")
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    doc_fill = PatternFill("solid", fgColor="7F6000")
    body_font = Font(name="Arial", size=10)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
        for c in ws[1]:
            c.font = head_font
            c.fill = doc_fill if ws.title in docs else head_fill
            c.alignment = Alignment(vertical="center")
            entry = None if ws.title in docs else lookup(ws.title, c.value)
            if entry:  # hover note on the header: meaning, unit, how it is used
                unit, meaning, use = entry
                text = meaning + (f"\nUnit: {unit}" if unit and unit != "-" else "") + (f"\n{use}" if use else "")
                note = Comment(text, "f1_to_excel")
                note.width, note.height = 320, 40 + 16 * (len(text) // 45 + text.count("\n") + 1)
                c.comment = note
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        if ws.title in docs:
            wide = {"README": {"A": 26, "B": 110},
                    "DataDictionary": {"A": 16, "B": 30, "C": 18, "D": 70, "E": 70},
                    "Legend": {"A": 28, "B": 24, "C": 110}}[ws.title]
            for letter, w in wide.items():
                ws.column_dimensions[letter].width = w
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")


def relabel_workbook(path):
    """Refresh README, DataDictionary, Legend and the header notes of an existing workbook,
    without re-downloading anything. Data sheets are left exactly as they are."""
    import openpyxl
    path = Path(path)
    old_readme = pd.read_excel(path, sheet_name="README")
    generated = dict(zip(old_readme["Item"], old_readme["Value"])).get("Generated (UTC)")
    races = pd.read_excel(path, sheet_name="Races")
    wb = openpyxl.load_workbook(path)
    for name in ("README", "DataDictionary", "Legend"):
        if name in wb.sheetnames:
            del wb[name]
    frames = {ws.title: pd.DataFrame(columns=[c.value for c in ws[1]]) for ws in wb.worksheets}
    check_dictionary(frames)
    docs = {
        "README": readme_frame(int(races["Year"].iloc[0]), not path.name.endswith("_dry_races.xlsx"),
                               races, generated=generated),
        "DataDictionary": dictionary_frame(frames),
        "Legend": pd.DataFrame(LEGEND, columns=["Group", "Code / term", "Meaning"]),
    }
    for name, df in docs.items():
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else v for v in row])
    wb._sheets.sort(key=lambda ws: SHEET_ORDER.index(ws.title) if ws.title in SHEET_ORDER else 99)
    wb.active = 0
    style_workbook(wb)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--years", type=int, nargs="+")
    ap.add_argument("--relabel", nargs="+", metavar="XLSX",
                    help="only refresh the legends (README, DataDictionary, Legend, header notes) of existing workbooks")
    ap.add_argument("--rounds", type=int, nargs="*", help="only these round numbers (for testing)")
    ap.add_argument("--include-wet", action="store_true", help="keep wet/mixed races")
    ap.add_argument("--out", default="f1_data", help="output folder for workbooks")
    ap.add_argument("--cache", default="fastf1_cache", help="FastF1 cache folder")
    ap.add_argument("--offline", action="store_true",
                    help="use the FastF1 cache only (no network, no F1TV token); uncached races show FAILED")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(message)s")
    if args.relabel:
        for f in args.relabel:
            relabel_workbook(f)
            log.info("relabelled %s", f)
        return
    if not args.years:
        ap.error("--years is required (or use --relabel)")
    logging.getLogger("fastf1").setLevel(logging.WARNING)
    out, cache = Path(args.out), Path(args.cache)
    out.mkdir(parents=True, exist_ok=True)
    cache.mkdir(parents=True, exist_ok=True)

    for year in args.years:
        races, sheets = load_season(year, args.rounds, args.include_wet, cache, offline=args.offline)
        sheets["Races"] = races
        sheets["README"] = readme_frame(year, args.include_wet, races)
        sheets["Legend"] = pd.DataFrame(LEGEND, columns=["Group", "Code / term", "Meaning"])
        check_dictionary(sheets)
        sheets["DataDictionary"] = dictionary_frame(sheets)
        suffix = "all_races" if args.include_wet else "dry_races"
        path = out / f"F1_{year}_{suffix}.xlsx"
        write_workbook(path, sheets)
        ok = races["LoadStatus"].eq("OK").sum() if "LoadStatus" in races else 0
        failed = races["LoadStatus"].astype(str).str.startswith("FAILED").sum()
        log.info("%d: wrote %s  (%d loaded, %d included, %d failed)", year, path, ok,
                 int(races["Included"].fillna(False).sum()), failed)
        if failed:
            log.info("  Some races failed (a rate limit, or a missing/expired F1TV token). Re-run the same command later; cached races are skipped quickly.")


if __name__ == "__main__":
    sys.exit(main())
